import re
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from operations.contracts import ServiceError
from operations.modules.master_data.application.contracts import DataSchema, RecordValues
from operations.modules.master_data.infrastructure.tabular import TabularFiles
from operations.modules.projects.application.contracts import LifecycleDefinition
from operations.modules.projects.domain.lifecycle import Lifecycle
from pydantic import ValidationError


def test_lifecycle_graph_requires_reachable_states_and_terminal_paths() -> None:
    assert Lifecycle().valid()
    assert Lifecycle().allows("planned", "active")
    assert not Lifecycle().allows("planned", "closed")
    for changes in (
        {"states": ("planned", "archived", "orphan"), "transitions": (("planned", "archived"),)},
        {"transitions": (("planned", "active"),)},
        {"states": ("planned", "archived", "planned")},
        {"transitions": (("archived", "planned"),)},
    ):
        with pytest.raises(ValidationError):
            LifecycleDefinition(**changes)


def test_declarative_fields_reject_duplicates_expressions_and_wrong_reference_contracts() -> None:
    for fields in (
        [{"key": "x", "kind": "python"}],
        [{"key": "x", "kind": "enum"}],
        [{"key": "x", "kind": "reference"}],
        [{"key": "x", "kind": "text", "choices": ["unexpected"]}],
        [{"key": "x", "kind": "text"}, {"key": "x", "kind": "integer"}],
    ):
        with pytest.raises(ValidationError):
            DataSchema.model_validate({"fields": fields})
    with pytest.raises(ValidationError):
        RecordValues.model_validate({"fields": {"x": 1.2}})


def test_formula_safe_csv_and_xlsx_export() -> None:
    files = TabularFiles()
    rows = [["id", "name"], ["stable-id", '=HYPERLINK("https://attacker.test")']]
    csv = files.export(rows, "csv")
    assert b"'=HYPERLINK" in csv
    workbook = load_workbook(BytesIO(files.export(rows, "xlsx")), data_only=False)
    try:
        assert workbook.active is not None
        cell = workbook.active["B2"]
        assert cell.data_type == "s" and cell.value == rows[1][1]
    finally:
        workbook.close()


def test_tabular_size_corruption_and_macro_boundaries() -> None:
    files = TabularFiles()
    for data, format in ((b"\xff", "csv"), (b"not-a-workbook", "xlsx")):
        with pytest.raises(ServiceError):
            files.parse(data, format)  # type: ignore[arg-type]
    with pytest.raises(ServiceError, match="import_too_large"):
        files.parse(b"a" * (4 * 1024 * 1024 + 1), "csv")
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("xl/vbaProject.bin", b"macro")
    with pytest.raises(ServiceError, match="unsafe_workbook"):
        files.parse(output.getvalue(), "xlsx")


def test_xlsx_import_reads_actual_cells_when_dimensions_underreport() -> None:
    workbook = Workbook()
    assert workbook.active is not None
    workbook.active.append(["code", "name"])
    workbook.active.append(["one", "First"])
    source = BytesIO()
    workbook.save(source)
    workbook.close()
    rewritten = BytesIO()
    with ZipFile(source) as archive, ZipFile(rewritten, "w", compression=ZIP_DEFLATED) as output:
        for entry in archive.infolist():
            data = archive.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', data)
            output.writestr(entry, data)
    assert TabularFiles().parse(rewritten.getvalue(), "xlsx") == [
        ["code", "name"],
        ["one", "First"],
    ]
