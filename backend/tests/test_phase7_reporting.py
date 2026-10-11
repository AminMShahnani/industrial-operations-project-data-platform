import csv
from datetime import UTC, datetime, timedelta
from io import BytesIO, StringIO
from types import SimpleNamespace
from typing import Any, cast
from uuid import UUID, uuid7

import pytest
from alembic.config import Config
from alembic.migration import MigrationContext
from alembic.operations import Operations
from alembic.script import ScriptDirectory
from openpyxl import load_workbook
from operations.composition import compose
from operations.modules.files.infrastructure.adapters import S3Storage
from operations.modules.identity.application.contracts import Principal, RequestContext
from operations.modules.identity.infrastructure.persistence import UserRow
from operations.modules.reporting.application.datasets import ReportingService
from operations.modules.reporting.infrastructure.rendering import TabularReportRenderer
from operations.modules.scheduling.application.contracts import Assignment
from operations.modules.tasks.application.contracts import Task
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError
from test_oidc import signing_key as signing_key
from test_phase1_api import Api
from test_phase1_api import api as api
from test_phase3_api import publish, setup
from test_phase4_api import activate, generate, schedule_setup


@pytest.mark.integration
def test_saved_report_migration_empty_roundtrip(api: Api) -> None:
    scripts = ScriptDirectory.from_config(Config("alembic.ini"))
    jobs = scripts.get_revision("92e3b0c4a7d1")
    recovery = scripts.get_revision("a16c3e902bf4")
    dashboards = scripts.get_revision("8d77c0f21a46")
    reports = scripts.get_revision("41a9c1d7e5b0")
    assert jobs and dashboards and reports and recovery
    with Operations.context(MigrationContext.configure(api.session.connection())):
        recovery.module.downgrade()
        jobs.module.downgrade()
        dashboards.module.downgrade()
        reports.module.downgrade()
        assert api.session.scalar(text("SELECT to_regclass('saved_report_versions')")) is None
        reports.module.upgrade()
        dashboards.module.upgrade()
        jobs.module.upgrade()
        recovery.module.upgrade()
        assert api.session.scalar(text("SELECT to_regclass('saved_report_versions')"))


@pytest.mark.integration
def test_export_jobs_are_private_reproducible_and_reauthorized_at_download(
    api: Api, monkeypatch: pytest.MonkeyPatch
) -> None:
    _, workspace_id, route, _ = setup(api)
    report_response = api.client.post(
        route + "/reports/saved",
        headers=api.headers("admin"),
        json={
            "name": "Private export",
            "definition": {
                "dataset": "overdue_tasks",
                "scope": "personal",
                "columns": ["id", "name"],
                "filters": [],
                "sort_by": "id",
                "sort_direction": "asc",
                "group_by": "none",
            },
        },
    )
    assert report_response.status_code == 201, report_response.text
    report_id = report_response.json()["id"]
    created = api.client.post(
        route + f"/reports/saved/{report_id}/export-jobs",
        headers=api.headers("admin"),
        json={"format": "pdf"},
    )
    assert created.status_code == 202, created.text
    job = created.json()
    assert job["state"] == "pending"
    assert job["report_version"] == 1
    assert job["report_digest"] == report_response.json()["definition_sha256"]
    assert job["expires_at"]
    org_id = UUID(route.split("/organizations/")[1].split("/")[0])
    token = api.invitation(org_id, workspace_id, role="Viewer")
    api.accept(org_id, token)
    assert (
        api.client.get(
            route + f"/reports/export-jobs/{job['id']}", headers=api.headers("invitee")
        ).status_code
        == 404
    )
    assert (
        api.client.get(
            route + f"/reports/export-jobs/{job['id']}/download", headers=api.headers("admin")
        ).status_code
        == 409
    )

    artifacts: dict[str, bytes] = {}

    def put(_storage: S3Storage, key: str, data: bytes, _content_type: str) -> None:
        artifacts[key] = data

    def get(_storage: S3Storage, key: str) -> bytes:
        return artifacts[key]

    monkeypatch.setattr(S3Storage, "put", put)
    monkeypatch.setattr(S3Storage, "get", get)
    principal = Principal("https://identity.example.test", "admin")
    context = RequestContext(principal, uuid7(), uuid7())
    services = compose(api.session, principal)
    rows = [
        SimpleNamespace(
            id=uuid7(),
            project_id=None,
            kind="generic",
            name=f"Task {index:03d}",
            state="open",
            due_at=datetime.now(UTC) - timedelta(minutes=2),
        )
        for index in range(101)
    ]

    def page_inbox(*args: Any, **_kwargs: Any) -> tuple[list[SimpleNamespace], UUID | None]:
        after = cast(UUID | None, args[7])
        selected = [item for item in rows if after is None or item.id > after]
        return selected[:100], selected[99].id if len(selected) > 100 else None

    monkeypatch.setattr(services.reporting.tasks, "inbox", page_inbox)
    export_jobs = services.reporting.export_jobs
    assert export_jobs is not None
    result = services.reporting.run_export_job(
        context, org_id, UUID(job["id"]), export_jobs, services.files.storage
    )
    assert result is not None and result.state == "completed"
    replayed = services.reporting.run_export_job(
        context, org_id, UUID(job["id"]), export_jobs, services.files.storage
    )
    assert replayed is not None and replayed.state == "completed"
    assert len(artifacts) == 1
    downloaded = api.client.get(
        route + f"/reports/export-jobs/{job['id']}/download", headers=api.headers("admin")
    )
    assert downloaded.status_code == 200, downloaded.text
    assert downloaded.content.startswith(b"%PDF-1.4")
    assert b"/Count 2" in downloaded.content
    assert b"Task 100" in downloaded.content
    assert downloaded.headers["cache-control"] == "private, no-store"
    revision = ScriptDirectory.from_config(Config("alembic.ini")).get_revision("92e3b0c4a7d1")
    assert revision
    with (
        Operations.context(MigrationContext.configure(api.session.connection())),
        pytest.raises(DBAPIError, match="Report export job history prevents rollback"),
        api.session.begin_nested(),
    ):
        revision.module.downgrade()


@pytest.mark.integration
def test_dashboard_migration_empty_roundtrip(api: Api) -> None:
    revision = ScriptDirectory.from_config(Config("alembic.ini")).get_revision("8d77c0f21a46")
    assert revision
    with Operations.context(MigrationContext.configure(api.session.connection())):
        revision.module.downgrade()
        assert api.session.scalar(text("SELECT to_regclass('dashboard_versions')")) is None
        revision.module.upgrade()
        assert api.session.scalar(text("SELECT to_regclass('dashboard_versions')"))


@pytest.mark.integration
def test_report_export_job_migration_empty_roundtrip(api: Api) -> None:
    scripts = ScriptDirectory.from_config(Config("alembic.ini"))
    revision = scripts.get_revision("92e3b0c4a7d1")
    recovery = scripts.get_revision("a16c3e902bf4")
    assert revision and recovery
    with Operations.context(MigrationContext.configure(api.session.connection())):
        recovery.module.downgrade()
        revision.module.downgrade()
        assert api.session.scalar(text("SELECT to_regclass('report_export_jobs')")) is None
        revision.module.upgrade()
        recovery.module.upgrade()
        assert api.session.scalar(text("SELECT to_regclass('report_export_jobs')"))


@pytest.mark.integration
def test_overdue_dataset_is_bounded_and_personal_scope_is_preserved(api: Api) -> None:
    route, schedule, _ = schedule_setup(api)
    activate(api, route, schedule)
    assert generate(api, route, schedule)["created"] == 1
    response = api.client.get(route + "/reports/overdue-tasks", headers=api.headers("admin"))
    assert response.status_code == 200, response.text
    page = response.json()
    assert len(page["items"]) == 1
    assert set(page["items"][0]) == {"id", "project_id", "kind", "name", "state", "due_at"}
    assert page["items"][0]["state"] == "open"
    assert page["next_cursor"] is None
    aggregate = api.client.get(
        route + "/reports/overdue-tasks?scope=aggregate", headers=api.headers("admin")
    )
    assert aggregate.status_code == 200, aggregate.text
    assert aggregate.json()["items"] == page["items"]

    org, workspace = route.split("/organizations/")[1].split("/workspaces/")
    token = api.invitation(UUID(org), UUID(workspace), role="Viewer")
    api.accept(UUID(org), token)
    viewer = api.client.get(route + "/reports/overdue-tasks", headers=api.headers("invitee"))
    assert viewer.status_code == 200, viewer.text
    assert viewer.json()["items"] == []
    denied = api.client.get(
        route + "/reports/overdue-tasks?scope=aggregate", headers=api.headers("invitee")
    )
    assert denied.status_code == 403


def test_overdue_dataset_uses_task_query_and_returns_minimal_projection() -> None:
    organization, workspace, user, identifier = uuid7(), uuid7(), uuid7(), uuid7()
    now = datetime.now(UTC)
    task = Task(
        id=identifier,
        organization_id=organization,
        workspace_id=workspace,
        project_id=None,
        kind="generic",
        origin_id=uuid7(),
        name="Check instrument",
        timezone="UTC",
        occurs_at=now,
        due_at=now,
        assignment=Assignment(kind="user", target_id=user),
        assignment_key="user:" + str(user),
        recipient_ids=[user],
        reminder_offsets=[],
        state="open",
    )
    cursor = uuid7()

    modes: list[str] = []

    class Tasks:
        def inbox(
            self,
            actor: RequestContext,
            org: UUID,
            selected_workspace: UUID,
            project: UUID | None,
            view: str,
            mode: str,
            target: UUID | None,
            after: UUID | None,
            filters: list[object] | None = None,
            sort_direction: str = "asc",
        ) -> tuple[list[Task], UUID | None]:
            assert actor.principal.subject == "reporter"
            assert (org, selected_workspace, project) == (organization, workspace, None)
            assert (view, target, after) == ("overdue", None, cursor)
            assert sort_direction == "asc"
            modes.append(mode)
            return [task], cursor

    service = object.__new__(ReportingService)
    untyped_service = cast(Any, service)
    untyped_service.tasks = Tasks()
    untyped_service.reports = None

    class Authorization:
        def require(self, actor: RequestContext, permission: str, scope: object) -> None:
            assert permission in {"task.read", "schedule.manage"}

    untyped_service.authorization = Authorization()
    untyped_service.workspaces = None
    untyped_service.projects = None
    untyped_service.audit = None
    result = service._read(
        RequestContext(Principal("issuer", "reporter"), uuid7(), uuid7()),
        organization,
        workspace,
        None,
        "personal",
        cursor,
    )
    assert result.next_cursor == cursor
    assert result.items[0].id == identifier
    assert result.items[0].name == "Check instrument"
    selected = service._read(
        RequestContext(Principal("issuer", "reporter"), uuid7(), uuid7()),
        organization,
        workspace,
        None,
        "aggregate",
        cursor,
        columns=["name"],
    )
    assert selected.items[0].id is None
    assert selected.items[0].name == "Check instrument"
    assert modes == ["personal", "project"]


def test_tabular_report_exports_treat_user_text_as_literal() -> None:
    renderer = TabularReportRenderer()
    csv_bytes = renderer.render(["name"], [["=1+1"]], "csv")
    assert list(csv.reader(StringIO(csv_bytes.decode("utf-8-sig")))) == [["name"], ["'=1+1"]]
    assert csv_bytes == renderer.render(["name"], [["=1+1"]], "csv")
    workbook_bytes = renderer.render(["name"], [["=1+1"]], "xlsx")
    assert workbook_bytes == renderer.render(["name"], [["=1+1"]], "xlsx")
    workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=False)
    sheet = workbook.active
    assert sheet is not None
    cell = sheet["A2"]
    assert cell.value == "=1+1"
    assert cell.data_type == "s"
    workbook.close()


def test_pdf_report_export_is_text_only_and_paginates_without_truncating() -> None:
    renderer = TabularReportRenderer()
    payload = renderer.render(
        ["name"], [[f"line {index} (literal)\\ safe"] for index in range(130)], "pdf"
    )
    assert payload.startswith(b"%PDF-1.4")
    assert b"/Count 3" in payload
    assert b"(line 129 \\(literal\\)\\\\ safe) Tj" in payload


@pytest.mark.integration
def test_saved_reports_are_private_versioned_and_reauthorize_runs(api: Api) -> None:
    org_id, workspace_id, route, form = setup(api)
    publish(api, route, form)
    user = api.session.scalar(
        select(UserRow).where(UserRow.organization_id == org_id, UserRow.subject == "admin")
    )
    assert user
    schedule_response = api.client.post(
        route + "/schedules",
        headers=api.headers("admin"),
        json={
            "name": "Shared operational work",
            "definition": {
                "form_id": form,
                "form_number": 1,
                "recurrence": {
                    "kind": "one_time",
                    "at": (datetime.now(UTC) - timedelta(hours=1)).isoformat(),
                },
                "assignments": [{"kind": "user", "target_id": str(user.id)}],
                "due_after_seconds": 60,
                "reminder_offsets": [],
            },
        },
    )
    assert schedule_response.status_code == 201, schedule_response.text
    schedule = schedule_response.json()["id"]
    activate(api, route, schedule)
    assert generate(api, route, schedule)["created"] == 1
    second_schedule_response = api.client.post(
        route + "/schedules",
        headers=api.headers("admin"),
        json={
            "name": "Second overdue operational work",
            "definition": {
                "form_id": form,
                "form_number": 1,
                "recurrence": {
                    "kind": "one_time",
                    "at": (datetime.now(UTC) - timedelta(hours=2)).isoformat(),
                },
                "assignments": [{"kind": "user", "target_id": str(user.id)}],
                "due_after_seconds": 60,
                "reminder_offsets": [],
            },
        },
    )
    assert second_schedule_response.status_code == 201, second_schedule_response.text
    second_schedule = second_schedule_response.json()["id"]
    activate(api, route, second_schedule)
    assert generate(api, route, second_schedule)["created"] == 1
    body = {
        "name": "My overdue tasks",
        "definition": {
            "dataset": "overdue_tasks",
            "scope": "personal",
            "columns": ["id", "name", "state", "due_at"],
            "filters": [{"field": "state", "value": "open"}],
            "sort_by": "id",
            "sort_direction": "desc",
            "group_by": "none",
        },
    }
    created = api.client.post(route + "/reports/saved", headers=api.headers("admin"), json=body)
    assert created.status_code == 201, created.text
    report = created.json()
    assert report["version"] == 1
    assert report["definition_sha256"]
    report_id = report["id"]
    owner = api.client.get(route + f"/reports/saved/{report_id}", headers=api.headers("admin"))
    assert owner.status_code == 200
    run = api.client.get(route + f"/reports/saved/{report_id}/run", headers=api.headers("admin"))
    assert run.status_code == 200, run.text
    result_ids = [UUID(item["id"]) for item in run.json()["items"]]
    assert len(result_ids) == 2
    assert result_ids == sorted(result_ids, reverse=True)
    csv_export = api.client.get(
        route + f"/reports/saved/{report_id}/export?format=csv", headers=api.headers("admin")
    )
    assert csv_export.status_code == 200, csv_export.text
    assert "attachment; filename=" in csv_export.headers["content-disposition"]
    assert csv_export.headers["cache-control"] == "private, no-store"
    csv_rows = list(csv.reader(StringIO(csv_export.content.decode("utf-8-sig"))))
    assert csv_rows[0] == ["id", "name", "state", "due_at"]
    assert len(csv_rows) == 3
    xlsx_export = api.client.get(
        route + f"/reports/saved/{report_id}/export?format=xlsx", headers=api.headers("admin")
    )
    assert xlsx_export.status_code == 200, xlsx_export.text
    exported_workbook = load_workbook(BytesIO(xlsx_export.content), read_only=True, data_only=False)
    exported_sheet = exported_workbook.active
    assert exported_sheet is not None
    values = list(exported_sheet.values)
    assert values[0] == ("id", "name", "state", "due_at")
    assert len(values) == 3
    exported_workbook.close()
    grouped_body = {
        **body,
        "name": "Overdue tasks by state",
        "definition": {
            **cast(dict[str, Any], body["definition"]),
            "columns": ["state"],
            "sort_direction": "asc",
            "group_by": "state",
        },
    }
    grouped_report = api.client.post(
        route + "/reports/saved", headers=api.headers("admin"), json=grouped_body
    )
    assert grouped_report.status_code == 201, grouped_report.text
    invalid_group_sort = api.client.post(
        route + "/reports/saved",
        headers=api.headers("admin"),
        json={
            **grouped_body,
            "name": "Invalid group sort",
            "definition": {
                **cast(dict[str, Any], grouped_body["definition"]),
                "sort_direction": "desc",
            },
        },
    )
    assert invalid_group_sort.status_code == 422
    grouped_result = api.client.get(
        route + f"/reports/saved/{grouped_report.json()['id']}/run",
        headers=api.headers("admin"),
    )
    assert grouped_result.status_code == 200, grouped_result.text
    assert grouped_result.json()["items"] == []
    assert grouped_result.json()["groups"] == [{"state": "open", "count": 2}]
    due_week_body = {
        **body,
        "name": "Current overdue counts by due week",
        "definition": {
            **cast(dict[str, Any], body["definition"]),
            "columns": ["due_at"],
            "sort_direction": "asc",
            "group_by": "due_week",
        },
    }
    due_week_report = api.client.post(
        route + "/reports/saved", headers=api.headers("admin"), json=due_week_body
    )
    assert due_week_report.status_code == 201, due_week_report.text
    due_week_id = due_week_report.json()["id"]
    due_week_result = api.client.get(
        route + f"/reports/saved/{due_week_id}/run", headers=api.headers("admin")
    )
    assert due_week_result.status_code == 200, due_week_result.text
    due_buckets = due_week_result.json()["due_week_groups"]
    assert len(due_buckets) == 13
    assert due_buckets[0]["week_start"] is None
    assert due_buckets[0]["label"].startswith("Before ")
    assert sum(item["count"] for item in due_buckets) == 2
    assert len([item for item in due_buckets if item["week_start"] is not None]) == 12
    due_week_export = api.client.get(
        route + f"/reports/saved/{due_week_id}/export?format=csv",
        headers=api.headers("admin"),
    )
    assert due_week_export.status_code == 200
    due_week_export_rows = list(csv.reader(StringIO(due_week_export.content.decode("utf-8-sig"))))
    assert due_week_export_rows[0] == ["due_week", "count"]
    assert len(due_week_export_rows) == 14
    second_due_week_report = api.client.post(
        route + "/reports/saved",
        headers=api.headers("admin"),
        json={**due_week_body, "name": "A second due-week view"},
    )
    assert second_due_week_report.status_code == 201, second_due_week_report.text
    trend_dashboard = api.client.post(
        route + "/reports/dashboards",
        headers=api.headers("admin"),
        json={
            "name": "Current due-week trend",
            "definition": {
                "widgets": [
                    {"report_id": due_week_id, "kind": "trend", "title": "Overdue trend"},
                    {
                        "report_id": second_due_week_report.json()["id"],
                        "kind": "line",
                        "title": "Overdue line",
                    },
                ]
            },
        },
    )
    assert trend_dashboard.status_code == 201, trend_dashboard.text
    trend_result = api.client.get(
        route + f"/reports/dashboards/{trend_dashboard.json()['id']}/run",
        headers=api.headers("admin"),
    )
    assert trend_result.status_code == 200, trend_result.text
    assert [item["kind"] for item in trend_result.json()["widgets"]] == ["trend", "line"]
    assert all(
        len(item["report"]["due_week_groups"]) == 13 for item in trend_result.json()["widgets"]
    )
    status_dashboard_body = {
        "name": "Open task status",
        "definition": {
            "widgets": [
                {
                    "report_id": grouped_report.json()["id"],
                    "kind": "status_distribution",
                    "title": "Open work by state",
                }
            ]
        },
    }
    status_dashboard = api.client.post(
        route + "/reports/dashboards",
        headers=api.headers("admin"),
        json=status_dashboard_body,
    )
    assert status_dashboard.status_code == 201, status_dashboard.text
    status_dashboard_run = api.client.get(
        route + f"/reports/dashboards/{status_dashboard.json()['id']}/run",
        headers=api.headers("admin"),
    )
    assert status_dashboard_run.status_code == 200, status_dashboard_run.text
    status_widget = status_dashboard_run.json()["widgets"][0]
    assert status_widget["kind"] == "status_distribution"
    assert status_widget["report"]["groups"] == [{"state": "open", "count": 2}]
    bar_dashboard = api.client.post(
        route + "/reports/dashboards",
        headers=api.headers("admin"),
        json={
            "name": "Open task chart",
            "definition": {
                "widgets": [
                    {
                        "report_id": grouped_report.json()["id"],
                        "kind": "bar",
                        "title": "Open work by state",
                    }
                ]
            },
        },
    )
    assert bar_dashboard.status_code == 201, bar_dashboard.text
    bar_result = api.client.get(
        route + f"/reports/dashboards/{bar_dashboard.json()['id']}/run",
        headers=api.headers("admin"),
    )
    assert bar_result.status_code == 200, bar_result.text
    assert bar_result.json()["widgets"][0]["kind"] == "bar"
    assert bar_result.json()["widgets"][0]["report"]["groups"] == [{"state": "open", "count": 2}]
    metric_dashboard = api.client.post(
        route + "/reports/dashboards",
        headers=api.headers("admin"),
        json={
            "name": "Overdue task metric",
            "definition": {
                "widgets": [
                    {
                        "report_id": grouped_report.json()["id"],
                        "kind": "metric",
                        "title": "Overdue tasks",
                    }
                ]
            },
        },
    )
    assert metric_dashboard.status_code == 201, metric_dashboard.text
    metric_result = api.client.get(
        route + f"/reports/dashboards/{metric_dashboard.json()['id']}/run",
        headers=api.headers("admin"),
    )
    assert metric_result.status_code == 200, metric_result.text
    assert metric_result.json()["widgets"][0]["metric_value"] == 2
    assert (
        api.client.post(
            route + "/reports/dashboards",
            headers=api.headers("admin"),
            json={
                "name": "Invalid status widget",
                "definition": {
                    "widgets": [
                        {
                            "report_id": report_id,
                            "kind": "status_distribution",
                            "title": "Not grouped",
                        }
                    ]
                },
            },
        ).status_code
        == 422
    )
    assert (
        api.client.post(
            route + "/reports/dashboards",
            headers=api.headers("admin"),
            json={
                "name": "Invalid row widget",
                "definition": {
                    "widgets": [
                        {
                            "report_id": grouped_report.json()["id"],
                            "kind": "table",
                            "title": "Grouped rows",
                        }
                    ]
                },
            },
        ).status_code
        == 422
    )

    org, workspace = route.split("/organizations/")[1].split("/workspaces/")
    token = api.invitation(UUID(org), UUID(workspace), role="WorkspaceOwner")
    api.accept(UUID(org), token)
    private = api.client.get(route + f"/reports/saved/{report_id}", headers=api.headers("invitee"))
    assert private.status_code == 404
    private_list = api.client.get(route + "/reports/saved", headers=api.headers("invitee"))
    assert private_list.status_code == 200
    assert private_list.json()["items"] == []

    updated = api.client.put(
        route + f"/reports/saved/{report_id}",
        headers=api.headers("admin"),
        json={**body, "name": "Renamed report", "expected_version": 1},
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["version"] == 2
    stale = api.client.put(
        route + f"/reports/saved/{report_id}",
        headers=api.headers("admin"),
        json={**body, "expected_version": 1},
    )
    assert stale.status_code == 409
    denied_run = api.client.get(
        route + f"/reports/saved/{report_id}/run", headers=api.headers("invitee")
    )
    assert denied_run.status_code == 404
    denied_export = api.client.get(
        route + f"/reports/saved/{report_id}/export?format=csv", headers=api.headers("invitee")
    )
    assert denied_export.status_code == 404
    dashboard_body = {
        "name": "My operational dashboard",
        "definition": {
            "widgets": [{"report_id": report_id, "kind": "table", "title": "Open work"}]
        },
    }
    created_dashboard = api.client.post(
        route + "/reports/dashboards", headers=api.headers("admin"), json=dashboard_body
    )
    assert created_dashboard.status_code == 201, created_dashboard.text
    dashboard_id = created_dashboard.json()["id"]
    dashboard_result = api.client.get(
        route + f"/reports/dashboards/{dashboard_id}/run", headers=api.headers("admin")
    )
    assert dashboard_result.status_code == 200, dashboard_result.text
    assert len(dashboard_result.json()["widgets"][0]["report"]["items"]) == 2
    assert (
        api.client.get(
            route + f"/reports/dashboards/{dashboard_id}", headers=api.headers("invitee")
        ).status_code
        == 404
    )
    assert (
        api.client.get(route + "/reports/dashboards", headers=api.headers("invitee")).json()[
            "items"
        ]
        == []
    )
    dashboard_update = api.client.put(
        route + f"/reports/dashboards/{dashboard_id}",
        headers=api.headers("admin"),
        json={
            **dashboard_body,
            "name": "Renamed dashboard",
            "definition": {
                "widgets": [{"report_id": report_id, "kind": "overdue_list", "title": "Due"}]
            },
            "expected_version": 1,
        },
    )
    assert dashboard_update.status_code == 200, dashboard_update.text
    assert dashboard_update.json()["version"] == 2
    assert (
        api.client.post(
            route + f"/reports/dashboards/{dashboard_id}/archive?expected_version=2",
            headers=api.headers("admin"),
        ).status_code
        == 204
    )
    archived = api.client.post(
        route + f"/reports/saved/{report_id}/archive?expected_version=2",
        headers=api.headers("admin"),
    )
    assert archived.status_code == 204, archived.text
    hidden = api.client.get(route + f"/reports/saved/{report_id}", headers=api.headers("admin"))
    assert hidden.status_code == 404
    revision = ScriptDirectory.from_config(Config("alembic.ini")).get_revision("41a9c1d7e5b0")
    assert revision
    with (
        Operations.context(MigrationContext.configure(api.session.connection())),
        pytest.raises(DBAPIError, match="Saved report history prevents rollback"),
    ):
        revision.module.downgrade()


@pytest.mark.integration
def test_dashboard_migration_refuses_populated_rollback(api: Api) -> None:
    organization_id, workspace_id, route, _ = setup(api)
    saved = api.client.post(
        route + "/reports/saved",
        headers=api.headers("admin"),
        json={
            "name": "Source report",
            "definition": {
                "dataset": "overdue_tasks",
                "scope": "personal",
                "columns": ["id", "name"],
                "filters": [],
                "sort_by": "id",
                "sort_direction": "asc",
                "group_by": "none",
            },
        },
    )
    assert saved.status_code == 201, saved.text
    dashboard = api.client.post(
        route + "/reports/dashboards",
        headers=api.headers("admin"),
        json={
            "name": "Retained dashboard",
            "definition": {
                "widgets": [{"report_id": saved.json()["id"], "kind": "table", "title": "Source"}]
            },
        },
    )
    assert dashboard.status_code == 201, dashboard.text
    revision = ScriptDirectory.from_config(Config("alembic.ini")).get_revision("8d77c0f21a46")
    assert revision
    with (
        Operations.context(MigrationContext.configure(api.session.connection())),
        pytest.raises(DBAPIError, match="Dashboard history prevents rollback"),
    ):
        revision.module.downgrade()


@pytest.mark.integration
@pytest.mark.parametrize(
    "table",
    [
        "saved_report_identities",
        "saved_report_versions",
        "saved_report_states",
        "dashboard_versions",
        "dashboard_states",
    ],
)
def test_reporting_history_rejects_truncate(api: Api, table: str) -> None:
    with pytest.raises(DBAPIError), api.session.begin_nested():
        api.session.execute(text(f"TRUNCATE TABLE {table}"))
