import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Literal
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from operations.contracts import ServiceError


class TabularFiles:
    def parse(self, data: bytes, format: Literal["csv", "xlsx"]) -> list[list[str]]:
        if len(data) > 4 * 1024 * 1024:
            raise ServiceError(413, "import_too_large")
        try:
            if format == "csv":
                reader = csv.reader(StringIO(data.decode("utf-8-sig")), strict=True)
                rows: list[list[str]] = []
                for row in reader:
                    if len(rows) >= 1001 or len(row) > 60 or any(len(cell) > 2000 for cell in row):
                        raise ServiceError(413, "import_too_large")
                    rows.append(row)
                return rows
            with ZipFile(BytesIO(data)) as archive:
                entries = archive.infolist()
                if (
                    len(entries) > 2000
                    or sum(entry.file_size for entry in entries) > 20 * 1024 * 1024
                ):
                    raise ServiceError(413, "import_too_large")
                if any(
                    "vba" in entry.filename.lower() or "externallinks" in entry.filename.lower()
                    for entry in entries
                ):
                    raise ServiceError(422, "unsafe_workbook")
            workbook = load_workbook(
                BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            try:
                if len(workbook.worksheets) != 1:
                    raise ServiceError(422, "single_import_sheet_required")
                sheet = workbook.worksheets[0]
                if not isinstance(sheet, ReadOnlyWorksheet):
                    raise ServiceError(422, "invalid_import_file")
                # Dimensions are producer-supplied metadata, not a trustworthy bound.
                sheet.reset_dimensions()
                rows = []
                for cells in sheet.iter_rows():
                    if len(rows) >= 1001 or len(cells) > 60:
                        raise ServiceError(413, "import_too_large")
                    row = []
                    for cell in cells:
                        if cell.data_type == "f":
                            raise ServiceError(422, "formula_cells_prohibited")
                        value = cell.value
                        if isinstance(value, datetime):
                            text = value.date().isoformat()
                        elif isinstance(value, date):
                            text = value.isoformat()
                        elif isinstance(value, bool):
                            text = "true" if value else "false"
                        else:
                            text = "" if value is None else str(value)
                        if len(text) > 2000:
                            raise ServiceError(413, "import_too_large")
                        row.append(text)
                    rows.append(row)
                return rows
            finally:
                workbook.close()
        except (
            UnicodeError,
            csv.Error,
            BadZipFile,
            InvalidFileException,
            ParseError,
            DefusedXmlException,
            ValueError,
            TypeError,
            KeyError,
        ) as error:
            raise ServiceError(422, "invalid_import_file") from error

    def export(self, rows: list[list[str]], format: Literal["csv", "xlsx"]) -> bytes:
        if format == "csv":
            output = StringIO(newline="")
            writer = csv.writer(output)
            for row in rows:
                writer.writerow(
                    [
                        "'" + cell
                        if cell.lstrip().startswith(("=", "+", "-", "@"))
                        or cell.startswith(("\t", "\r"))
                        else cell
                        for cell in row
                    ]
                )
            return output.getvalue().encode("utf-8-sig")
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Master data")
        for row in rows:
            cells = []
            for text in row:
                cell = WriteOnlyCell(sheet, value=text)
                cell.data_type = "s"
                cells.append(cell)
            sheet.append(cells)
        output_bytes = BytesIO()
        workbook.save(output_bytes)
        workbook.close()
        return output_bytes.getvalue()
